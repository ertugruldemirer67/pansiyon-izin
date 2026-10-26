"""Excel'den öğrenci / veli listesi içe aktarma."""
import io
import re
import secrets
from typing import BinaryIO

from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from .models import Kullanici, Ogrenci, VeliOgrenci
from .security import kayit_kodu_uret, kod_hashle

GEREKLI = ("okul_no", "ogrenci_ad_soyad", "veli_ad_soyad", "veli_telefon")
ISTEGE_BAGLI = ("sinif", "oda", "yakinlik")


def telefon_normalle(deger) -> str:
    rakamlar = re.sub(r"\D", "", str(deger or ""))
    return rakamlar[-10:]


def _hucre(deger) -> str:
    if deger is None:
        return ""
    if isinstance(deger, float) and deger.is_integer():
        deger = int(deger)
    return str(deger).strip()


def ice_aktar(db: Session, dosya: BinaryIO) -> dict:
    """Satırları işler. Yeni oluşturulan hesaplar için tek seferlik kayıt kodlarını döndürür."""
    sayfa = load_workbook(dosya, read_only=True, data_only=True).active
    satirlar = sayfa.iter_rows(values_only=True)
    try:
        baslik = [_hucre(h).lower() for h in next(satirlar)]
    except StopIteration:
        return {"hata": "Dosya boş"}

    eksik = [k for k in GEREKLI if k not in baslik]
    if eksik:
        return {"hata": f"Eksik sütunlar: {', '.join(eksik)}"}
    idx = {k: baslik.index(k) for k in GEREKLI + ISTEGE_BAGLI if k in baslik}

    def al(satir, anahtar):
        i = idx.get(anahtar)
        return _hucre(satir[i]) if i is not None and i < len(satir) else ""

    rapor = {"ogrenci_yeni": 0, "ogrenci_guncel": 0, "veli_yeni": 0, "hatalar": [], "kodlar": []}

    for no, satir in enumerate(satirlar, start=2):
        if not any(satir):
            continue
        okul_no, ogr_ad = al(satir, "okul_no"), al(satir, "ogrenci_ad_soyad")
        veli_ad, tel = al(satir, "veli_ad_soyad"), telefon_normalle(al(satir, "veli_telefon"))
        if not (okul_no and ogr_ad and veli_ad) or len(tel) != 10:
            rapor["hatalar"].append(f"Satır {no}: eksik veya hatalı bilgi")
            continue

        ogrenci = db.scalars(select(Ogrenci).where(Ogrenci.okul_no == okul_no)).first()
        if ogrenci is None:
            ogr_kul = Kullanici(kullanici_adi=okul_no, ad_soyad=ogr_ad, rol="ogrenci",
                                aktif=False)
            kod = kayit_kodu_uret()
            ogr_kul.kayit_kodu_hash = kod_hashle(kod)
            if db.scalars(select(Kullanici).where(Kullanici.kullanici_adi == okul_no)).first():
                rapor["hatalar"].append(f"Satır {no}: {okul_no} kullanıcı adı başka hesapta kullanılıyor")
                continue
            db.add(ogr_kul)
            db.flush()
            ogrenci = Ogrenci(kullanici_id=ogr_kul.id, okul_no=okul_no, ad_soyad=ogr_ad,
                              sinif=al(satir, "sinif"), oda=al(satir, "oda"),
                              qr_token=secrets.token_urlsafe(16))
            db.add(ogrenci)
            rapor["ogrenci_yeni"] += 1
            rapor["kodlar"].append({"kullanici_adi": okul_no, "ad_soyad": ogr_ad,
                                    "rol": "ogrenci", "kod": kod})
        else:
            ogrenci.ad_soyad = ogr_ad
            ogrenci.sinif = al(satir, "sinif") or ogrenci.sinif
            ogrenci.oda = al(satir, "oda") or ogrenci.oda
            rapor["ogrenci_guncel"] += 1

        veli = db.scalars(select(Kullanici).where(Kullanici.kullanici_adi == tel)).first()
        if veli is None:
            kod = kayit_kodu_uret()
            veli = Kullanici(kullanici_adi=tel, ad_soyad=veli_ad, rol="veli", aktif=False,
                             kayit_kodu_hash=kod_hashle(kod))
            db.add(veli)
            db.flush()
            rapor["veli_yeni"] += 1
            rapor["kodlar"].append({"kullanici_adi": tel, "ad_soyad": veli_ad,
                                    "rol": "veli", "kod": kod})
        elif veli.rol != "veli":
            rapor["hatalar"].append(f"Satır {no}: {tel} numarası veli olmayan bir hesapta")
            continue

        bag = db.get(VeliOgrenci, (veli.id, ogrenci.id))
        if bag is None:
            db.add(VeliOgrenci(veli_id=veli.id, ogrenci_id=ogrenci.id,
                               yakinlik=al(satir, "yakinlik")))
        db.flush()

    return rapor


def sablon_xlsx() -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Liste"
    ws.append(["okul_no", "ogrenci_ad_soyad", "sinif", "oda", "veli_ad_soyad", "veli_telefon", "yakinlik"])
    ws.append(["1001", "Örnek Öğrenci", "9-A", "12", "Örnek Veli", "05321234567", "Anne"])
    for harf, genislik in zip("ABCDEFG", (10, 24, 8, 8, 24, 16, 10)):
        ws.column_dimensions[harf].width = genislik
    cikti = io.BytesIO()
    wb.save(cikti)
    return cikti.getvalue()

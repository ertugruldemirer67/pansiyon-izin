from datetime import datetime, timedelta

from tests.conftest import SIMDI
from tests.test_akis import talep_govdesi

# SIMDI: 6 Kasım 2026 Cuma 12:00
CUMA_17 = datetime(2026, 11, 6, 17, 0)
PAZAR_19 = datetime(2026, 11, 8, 19, 0)
PAZAR_21 = datetime(2026, 11, 8, 21, 0)
SALI_10 = datetime(2026, 11, 10, 10, 0)

KURAL = {"evci_pencereleri": [
    {"cikis_gun": 4, "cikis_saat": "17:00", "donus_gun": 6, "donus_saat": "20:00"}]}


def evci(oid, cikis, donus, **ek):
    return talep_govdesi(oid, cikis_zamani=cikis.isoformat(), donus_zamani=donus.isoformat(), **ek)


def test_ayarlar_varsayilan_ve_yetki(istemci):
    a = istemci("veli1").get("/ayarlar").json()
    assert a["evci_pencereleri"] == [] and a["aylik_evci_hakki"] == 0
    assert a["gunler"][4] == "Cuma"
    assert istemci("veli1").put("/ayarlar", json={"aylik_evci_hakki": 2}).status_code == 403
    assert istemci("belletmen1").put("/ayarlar", json={"aylik_evci_hakki": 2}).status_code == 403


def test_ayar_dogrulama(istemci):
    yon = istemci("yonetici1")
    assert yon.put("/ayarlar", json={"aylik_evci_hakki": -1}).status_code == 422
    assert yon.put("/ayarlar", json={"bilinmeyen": 1}).status_code == 422
    kotu = {"evci_pencereleri": [{"cikis_gun": 9, "cikis_saat": "17:00",
                                  "donus_gun": 6, "donus_saat": "20:00"}]}
    assert yon.put("/ayarlar", json=kotu).status_code == 422
    assert yon.put("/ayarlar", json={"evci_pencereleri": [
        {"cikis_gun": 4, "cikis_saat": "25:00", "donus_gun": 6, "donus_saat": "20:00"}]}
    ).status_code == 422


def test_evci_penceresi(istemci, veri):
    istemci("yonetici1").put("/ayarlar", json=KURAL)
    veli, oid = istemci("veli1"), veri["ogrenci"].id

    ok = veli.post("/talepler", json=evci(oid, CUMA_17, PAZAR_19))
    assert ok.status_code == 201
    veli.post(f"/talepler/{ok.json()['id']}/iptal")

    gec = veli.post("/talepler", json=evci(oid, CUMA_17, PAZAR_21))
    assert gec.status_code == 422 and "Cuma 17:00" in gec.json()["detail"]
    hafta_ici = veli.post("/talepler", json=evci(oid, SALI_10, SALI_10 + timedelta(hours=4)))
    assert hafta_ici.status_code == 422

    # evci dışındaki türler pencereden etkilenmez
    mazeret = veli.post("/talepler", json=evci(oid, SALI_10, SALI_10 + timedelta(hours=4),
                                               tur="mazeret"))
    assert mazeret.status_code == 201


def test_yonetici_acil_izin_kurali_asar(istemci, veri):
    istemci("yonetici1").put("/ayarlar", json=KURAL)
    r = istemci("yonetici1").post("/talepler", json=evci(
        veri["ogrenci"].id, SALI_10, SALI_10 + timedelta(hours=4), gerekce="Acil"))
    assert r.status_code == 201


def test_aylik_hak(istemci, veri):
    yon, veli, oid = istemci("yonetici1"), istemci("veli1"), veri["ogrenci"].id
    yon.put("/ayarlar", json={"aylik_evci_hakki": 2})
    assert veli.get(f"/ogrenciler/{oid}/evci-hakki").json() == {"kullanilan": 0, "hak": 2}

    g1 = datetime(2026, 11, 13, 17, 0)
    g2 = datetime(2026, 11, 20, 17, 0)
    g3 = datetime(2026, 11, 27, 17, 0)
    ids = []
    for g in (g1, g2):
        r = veli.post("/talepler", json=evci(oid, g, g + timedelta(days=2)))
        assert r.status_code == 201
        ids.append(r.json()["id"])
    r = veli.post("/talepler", json=evci(oid, g3, g3 + timedelta(days=2)))
    assert r.status_code == 422 and "hak" in r.json()["detail"]

    # iptal hak iade eder
    veli.post(f"/talepler/{ids[0]}/iptal")
    assert veli.post("/talepler", json=evci(oid, g3, g3 + timedelta(days=2))).status_code == 201

    # başka ay etkilenmez
    aralik = datetime(2026, 12, 4, 17, 0)
    assert veli.post("/talepler", json=evci(oid, aralik, aralik + timedelta(days=2))).status_code == 201


def test_tolerans_ayari(istemci, veri, saat):
    yon, veli, bel = istemci("yonetici1"), istemci("veli1"), istemci("belletmen1")
    yon.put("/ayarlar", json={"tolerans_dakika": 0})
    tid = veli.post("/talepler", json=talep_govdesi(veri["ogrenci"].id)).json()["id"]
    yon.post(f"/talepler/{tid}/yonetici-karar", json={"onay": True})
    saat.t = SIMDI + timedelta(hours=2) - timedelta(minutes=5)  # çıkıştan 5 dk önce
    assert bel.post("/hareket/okut", json={"qr_token": "qr-ali"}).status_code == 409
    yon.put("/ayarlar", json={"tolerans_dakika": 30})
    assert bel.post("/hareket/okut", json={"qr_token": "qr-ali"}).json()["tip"] == "cikis"


def test_okul_no_ile_okutma(istemci, veri, saat):
    yon, veli, bel = istemci("yonetici1"), istemci("veli1"), istemci("belletmen1")
    tid = veli.post("/talepler", json=talep_govdesi(veri["ogrenci"].id)).json()["id"]
    yon.post(f"/talepler/{tid}/yonetici-karar", json={"onay": True})
    saat.t = SIMDI + timedelta(hours=2)
    assert bel.post("/hareket/okut", json={}).status_code == 422
    assert bel.post("/hareket/okut", json={"okul_no": "9999"}).status_code == 404
    assert bel.post("/hareket/okut", json={"okul_no": "1001"}).json()["tip"] == "cikis"


def test_toplu_karar(istemci, veri):
    veli, yon = istemci("veli1"), istemci("yonetici1")
    oid = veri["ogrenci"].id
    t1 = veli.post("/talepler", json=talep_govdesi(oid)).json()["id"]
    g = datetime(2026, 12, 11, 17, 0)
    t2 = veli.post("/talepler", json=evci(oid, g, g + timedelta(days=2))).json()["id"]
    r = yon.post("/talepler/toplu-karar", json={"idler": [t1, t2, 999], "onay": True}).json()
    assert r["islenen"] == 2 and r["atlanan"][0]["id"] == 999
    assert istemci("belletmen1").post("/talepler/toplu-karar",
                                      json={"idler": [t1], "onay": True}).status_code == 403


def test_ogrenci_arama_yetkisi(istemci):
    assert [o["ad_soyad"] for o in istemci("belletmen1").get("/ogrenciler?q=ali").json()] == ["Ali Veli"]
    assert istemci("veli1").get("/ogrenciler").status_code == 403
    assert istemci("1001").get("/ogrenciler").status_code == 403


def test_personel_ve_kod_yenileme(uygulama, veri):
    from fastapi.testclient import TestClient
    from app.models import Kullanici
    from app.security import sifre_hashle
    from tests.conftest import SIFRE
    with uygulama.state.oturum_fabrikasi() as db:
        db.add(Kullanici(kullanici_adi="sysadmin", ad_soyad="Sys", rol="sistem_yoneticisi",
                         sifre_hash=sifre_hashle(SIFRE)))
        db.commit()
    admin = TestClient(uygulama)
    admin.post("/auth/giris", json={"kullanici_adi": "sysadmin", "sifre": SIFRE})

    assert admin.post("/admin/personel", json={"kullanici_adi": "bel2", "ad_soyad": "Yeni Belletmen",
                                               "rol": "veli"}).status_code == 422
    r = admin.post("/admin/personel", json={"kullanici_adi": "bel2", "ad_soyad": "Yeni Belletmen",
                                            "rol": "belletmen"})
    assert r.status_code == 201
    kod = r.json()["kod"]
    yeni = TestClient(uygulama)
    assert yeni.post("/auth/ilk-kayit", json={"kullanici_adi": "bel2", "kod": kod,
                                              "sifre": "belletmen-sifre"}).status_code == 200
    assert yeni.post("/auth/giris", json={"kullanici_adi": "bel2",
                                          "sifre": "belletmen-sifre"}).status_code == 200

    # şifre unutuldu: kod yenile, eski şifre geçersiz
    uid = r.json()["id"]
    yeni_kod = admin.post(f"/admin/kullanicilar/{uid}/kod-yenile").json()["kod"]
    assert TestClient(uygulama).post("/auth/giris", json={
        "kullanici_adi": "bel2", "sifre": "belletmen-sifre"}).status_code == 401
    assert TestClient(uygulama).post("/auth/ilk-kayit", json={
        "kullanici_adi": "bel2", "kod": yeni_kod, "sifre": "ikinci-sifre-1"}).status_code == 200
    assert admin.post(f"/admin/kullanicilar/{uid}/pasif").json() == {"tamam": True}
    assert TestClient(uygulama).post("/auth/giris", json={
        "kullanici_adi": "bel2", "sifre": "ikinci-sifre-1"}).status_code == 401


def test_statik_dosyalar(uygulama):
    from fastapi.testclient import TestClient
    c = TestClient(uygulama)
    assert c.get("/").status_code == 200
    assert c.get("/manifest.webmanifest").json()["display"] == "standalone"
    assert c.get("/sw.js").status_code == 200
    assert c.get("/bos-form").status_code == 200


def test_cifte_okutma_korumasi(istemci, veri, saat):
    yon, veli, bel = istemci("yonetici1"), istemci("veli1"), istemci("belletmen1")
    tid = veli.post("/talepler", json=talep_govdesi(veri["ogrenci"].id)).json()["id"]
    yon.post(f"/talepler/{tid}/yonetici-karar", json={"onay": True})
    saat.t = SIMDI + timedelta(hours=2)
    assert bel.post("/hareket/okut", json={"qr_token": "qr-ali"}).json()["tip"] == "cikis"
    # aynı kod hemen tekrar okutulursa dönüş sayılmaz
    r = bel.post("/hareket/okut", json={"qr_token": "qr-ali"})
    assert r.status_code == 409 and "az önce" in r.json()["detail"]
    saat.t += timedelta(minutes=2)
    assert bel.post("/hareket/okut", json={"qr_token": "qr-ali"}).json()["tip"] == "giris"


def test_sablon_indirme(uygulama, veri):
    import io
    from fastapi.testclient import TestClient
    from openpyxl import load_workbook
    from app.models import Kullanici
    from app.security import sifre_hashle
    from tests.conftest import SIFRE
    with uygulama.state.oturum_fabrikasi() as db:
        db.add(Kullanici(kullanici_adi="sysadmin", ad_soyad="Sys", rol="sistem_yoneticisi",
                         sifre_hash=sifre_hashle(SIFRE)))
        db.commit()
    admin = TestClient(uygulama)
    admin.post("/auth/giris", json={"kullanici_adi": "sysadmin", "sifre": SIFRE})
    r = admin.get("/admin/sablon")
    assert r.status_code == 200
    satirlar = list(load_workbook(io.BytesIO(r.content)).active.iter_rows(values_only=True))
    assert satirlar[0][:4] == ("okul_no", "ogrenci_ad_soyad", "sinif", "oda")
    # şablonun kendisi içe aktarılabilir
    r = admin.post("/admin/ice-aktar", files={"dosya": ("s.xlsx", io.BytesIO(
        admin.get("/admin/sablon").content))})
    assert r.status_code == 200 and r.json()["hatalar"] == []
    assert r.json()["ogrenci_yeni"] + r.json()["ogrenci_guncel"] == 1


def test_guvenlik_basliklari_ve_onbellek(uygulama):
    from fastapi.testclient import TestClient
    c = TestClient(uygulama)
    r = c.get("/")
    assert r.headers["x-frame-options"] == "DENY"
    assert "default-src 'self'" in r.headers["content-security-policy"]
    assert c.get("/saglik").headers["cache-control"] == "no-store"
